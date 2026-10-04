"""T17 et T18 adaptés : profils logiciels, zéro participant humain.
Les temps sont des temps d'automatisation, pas d'utilisation autonome.
"""
import os
import sys
import json
import time
import statistics
from pathlib import Path
os.environ['OS4_RUN']='sessions'
sys.path.insert(0,'/app/scripts')
import os4_campaign as c
from openpyxl import Workbook,load_workbook
OUT=Path('/app/os4/sessions');OUT.mkdir(exist_ok=True)
PROFILES=[('S01','agent terrain A'),('S02','agent terrain B'),('S03','gestionnaire A'),('S04','gestionnaire B'),('S05','superviseur MEAL'),('S06','direction')]
def ensure(response,statuses):
    if response.status_code not in statuses:raise RuntimeError(str(response.status_code)+' '+str(response.data))
    return response
def complete(d,project):
    path=c.root(d);c.prepare(d)
    ensure(c.call('R2','POST',path+'/notes',{'text':'FICTIF note interne'}),[201])
    reply=ensure(c.call('R2','POST',path+'/reponses',{'text':'FICTIF calendrier confirmé'}),[201]).data['id']
    ensure(c.call('R3','POST',path+f'/reponses/{reply}/validation',{'accepted':True}),[200])
    ensure(c.call('R3','POST',path+'/transitions',{'state':'CLOTURE','reason':'RESOLU'}),[200])
    return d
def main():
    results=[];comp=[]
    for pid,profile in PROFILES:
        c.RUN=pid;c.users.clear();c.tokens.clear();c.setup();d=None
        for task in range(1,7):
            start=time.perf_counter();status='Réussi';error=None
            try:
                if task==1:
                    rr=ensure(c.call('R0','POST','/api/v1/public/retours',{'key':pid+'-U1','text':'FICTIF dépôt anonyme '+pid},c.APIClient(REMOTE_ADDR='198.18.10.'+str(int(pid[1:])))),[201]);assert rr.data['reference']
                elif task==2:
                    rr=ensure(c.call('R1','POST','/api/v1/agent/sync',{'items':[{'key':pid+'-U2','text':'FICTIF saisie assistée','address':pid+'-CONTACT','consent':True}]}),[200]);assert rr.data['results'][0]['reference']
                elif task==3:
                    d=c.fresh(pid+'-U3');c.prepare(d);d.refresh_from_db();assert d.owner_id==c.users['R2'].pk and d.due
                elif task==4:
                    path=c.root(d);rid=ensure(c.call('R2','POST',path+'/reponses',{'text':'FICTIF réponse '+pid}),[201]).data['id'];ensure(c.call('R3','POST',path+f'/reponses/{rid}/validation',{'accepted':True}),[200]);d.refresh_from_db();assert d.state=='REPONSE_TRANSMISE'
                elif task==5:
                    ensure(c.call('R3','POST',c.root(d)+'/transitions',{'state':'CLOTURE','reason':'RESOLU'}),[200]);ensure(c.call('R3','POST',c.root(d)+'/transitions',{'state':'EN_TRAITEMENT','reason':'FICTIF retour '+pid}),[200]);d.refresh_from_db();assert d.reopen_count==1
                elif task==6:
                    project='OS4 SESSION '+pid
                    a=ensure(c.call('R3','POST','/api/v1/actions',{'title':'FICTIF amélioration '+pid,'project':project,'owner':c.users['R3'].pk,'due':'2026-12-01','dossiers':[str(d.pk)]}),[201]).data['id'];ensure(c.call('R3','PATCH',f'/api/v1/actions/{a}',{'status':'REALISEE','proof':'FICTIF compte rendu'}),[200]);ensure(c.call('R5','POST',f'/api/v1/actions/{a}/restitutions',{'date':'2026-10-04','channel':'Réunion fictive','zone':'Zone A','message':'FICTIF restitution'}),[201]);rr=c.call('R5','GET','/api/v1/indicateurs',{'project':project});assert rr.data['indicators']['I07']==100
            except Exception as ex:status='Échoué';error=str(ex)
            results.append({'profile':pid,'role_type':profile,'task':'U'+str(task),'status':status,'seconds':round(time.perf_counter()-start,4),'error':error,'human':False})
        # Twenty distinct comparison cases per software profile, never used to tune detector.
        register=Workbook();sheet=register.active;sheet.title='Registre';sheet.append(['ID','Réception','Texte','Type','Thème','Titulaire','Réponse','Statut','Motif','Action','Traçabilité'])
        order=['Excel','Écoute'] if int(pid[1:])%2 else ['Écoute','Excel']
        for tool in order:
            for n in range(20):
                start=time.perf_counter();key=f'{pid}-CMP-{n:02d}';status='Réussi';err=None
                try:
                    if tool=='Excel':
                        sheet.append([key,c.timezone.now().isoformat(),'FICTIF calendrier de distribution','PLAINTE','Distribution','Gestionnaire fictif','FICTIF calendrier confirmé','CLOTURE','RESOLU','Sans action','Réception qualification attribution réponse clôture'])
                        register.save(OUT/(pid+'-registre.xlsx'))
                        check=load_workbook(OUT/(pid+'-registre.xlsx'),read_only=True);assert check['Registre'].max_row==n+2;check.close()
                    else:
                        d=c.fresh(key);complete(d,'OS4 COMPARAISON');d.refresh_from_db();assert d.state=='CLOTURE' and d.responded and d.records.exists()
                except Exception as ex:status='Échoué';err=str(ex)
                comp.append({'profile':pid,'tool':tool,'order':order,'case':key,'seconds':round(time.perf_counter()-start,4),'status':status,'error':err,'human':False})
        print(pid,'6 tâches et 20 cas par outil terminés',flush=True)
    (OUT/'tasks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));(OUT/'comparison.json').write_text(json.dumps(comp,ensure_ascii=False,indent=2))
    summary={'software_profiles':6,'human_participants':0,'tasks_attempted':len(results),'tasks_passed':sum(x['status']=='Réussi' for x in results),'SUS':None,'K08_human_gain':None,'K09_human_success':None,'comparison_machine_only':{tool:{'attempts':sum(x['tool']==tool for x in comp),'passed':sum(x['tool']==tool and x['status']=='Réussi' for x in comp),'median_seconds':statistics.median(x['seconds'] for x in comp if x['tool']==tool)} for tool in ['Excel','Écoute']},'limitation':'Excel programmé et APIClient PostgreSQL sans interaction humaine ; temps non comparables au temps utilisateur ni au gain K08 ; pas de pause cognitive ni apprentissage mesuré'}
    (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));print(json.dumps(summary,ensure_ascii=False),flush=True)
if __name__=='__main__':main()
